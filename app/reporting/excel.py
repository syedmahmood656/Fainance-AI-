"""Excel workbook generation (OpenPyXL). Takes plain dict rows that came from
PostgreSQL; every number in the workbook is computed in Python from those rows."""
from collections import OrderedDict
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="3B6EA5")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MONEY = "#,##0.00"


def safe_cell(value: Any) -> Any:
    """Stop spreadsheet formula injection: a vendor named '=HYPERLINK(...)' from a
    malicious receipt must stay plain text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _header(ws, headers: List[str]) -> None:
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = f"A{ws.max_row + 1}"


def _autosize(ws) -> None:
    for col in ws.columns:
        longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, longest + 2), 60)


def build_workbook(expenses: List[Dict[str, Any]], period_label: str, currency: str,
                   category: Optional[str] = None, other_currency_count: int = 0) -> Workbook:
    """expenses: dicts with the final values (see tools.expense_to_dict)."""
    wb = Workbook()
    total = round(sum(e["expense_amount"] or 0 for e in expenses), 2)
    count = len(expenses)

    # ---- Summary ----
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Expense Report"
    ws["A1"].font = Font(bold=True, size=16)
    rows = [
        ("Period", period_label),
        ("Category filter", category or "All categories"),
        ("Currency", currency),
        ("Total expenses", total),
        ("Number of expenses", count),
        ("Average expense", round(total / count, 2) if count else 0),
        ("Corrected by a human", sum(1 for e in expenses if e.get("corrected_by_user"))),
        ("Generated at (UTC)", datetime.utcnow().strftime("%Y-%m-%d %H:%M")),
        ("Scope", "Confirmed expenses only"),
    ]
    if other_currency_count:
        rows.append(("Excluded (other currencies)", other_currency_count))
    for i, (label, value) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=safe_cell(value))
        if label in ("Total expenses", "Average expense"):
            cell.number_format = _MONEY
    _autosize(ws)

    # ---- Expenses (the actual filtered rows) ----
    ws = wb.create_sheet("Expenses")
    _header(ws, ["ID", "Date", "Vendor", "Invoice #", "Category", "Item", "Type", "Amount", "Currency",
                 "Subtotal", "Tax", "Description", "Corrected by user"])
    for e in expenses:
        ws.append([
            e["id"], e["expense_date"], safe_cell(e["vendor"]), safe_cell(e["invoice_number"]),
            e["category"], safe_cell(e["item"]), e["goods_or_service"], e["expense_amount"], e["currency"],
            e["subtotal"], e["tax"], safe_cell(e["description"]), "Yes" if e.get("corrected_by_user") else "No",
        ])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "DD MMM YYYY"
        for idx in (7, 9, 10):
            row[idx].number_format = _MONEY
    if count:
        ws.append([])
        ws.append(["", "", "", "", "", "", "TOTAL", total])
        ws.cell(row=ws.max_row, column=7).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=8).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=8).number_format = _MONEY
    _autosize(ws)

    # ---- Category Summary ----
    ws = wb.create_sheet("Category Summary")
    _header(ws, ["Category", "Total", "Count", "Share %"])
    by_cat: "OrderedDict[str, List[float]]" = OrderedDict()
    for e in expenses:
        slot = by_cat.setdefault(e["category"] or "Uncategorized", [0.0, 0])
        slot[0] += e["expense_amount"] or 0
        slot[1] += 1
    for name, (amt, n) in sorted(by_cat.items(), key=lambda kv: -kv[1][0]):
        ws.append([name, round(amt, 2), n, round(amt / total * 100, 1) if total else 0])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = _MONEY
    _autosize(ws)

    # ---- Monthly Summary ----
    ws = wb.create_sheet("Monthly Summary")
    _header(ws, ["Month", "Total", "Count"])
    by_month: Dict[str, List[float]] = {}
    for e in expenses:
        d = e["expense_date"]
        key = d.strftime("%Y-%m") if isinstance(d, (date, datetime)) else "Unknown"
        slot = by_month.setdefault(key, [0.0, 0])
        slot[0] += e["expense_amount"] or 0
        slot[1] += 1
    for key in sorted(by_month):
        ws.append([key, round(by_month[key][0], 2), by_month[key][1]])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = _MONEY
    _autosize(ws)
    return wb
