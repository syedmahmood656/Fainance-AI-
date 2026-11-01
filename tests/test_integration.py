"""End-to-end tests on SQLite with a fake LLM: receipt -> confirm/correct -> reports -> Excel -> receipt image."""
import asyncio
import io
import os
from datetime import date

import pytest
from openpyxl import load_workbook
from PIL import Image

from app.agent import graph as graph_module
from app.agent.graph import run_agent
from app.schemas.enums import ALL_CATEGORIES
from app.ai.llm_provider import LLMCallInfo, LLMUnavailableError
from app.database.models import Expense, ExpenseCorrection, ManualDraft, PendingEdit, SessionLocal, init_db
from app.reporting import tools
from app.schemas.expense import CategorizationResult, Period, ReceiptExtraction, ReportRequest
from app.services.expense_service import expense_service
from app.services.storage_service import storage_service

USER = 111
# Derived from the live category list, so editing your categories never breaks these tests.
CAT = ALL_CATEGORIES[1].value          # what the fake AI predicts
CAT_IDX_FIXED = 7                      # the category the human picks with the button
CAT_FIXED = ALL_CATEGORIES[CAT_IDX_FIXED].value


@pytest.fixture(autouse=True)
def fresh_db():
    init_db()
    with SessionLocal() as s:
        for model in (ExpenseCorrection, PendingEdit, ManualDraft, Expense):
            s.query(model).delete()
        s.commit()
    yield


def png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (200, 300), (240, 240, 240)).save(buf, "PNG")
    return buf.getvalue()


def agent(**state):
    base = {"user_id": USER, "input_type": "text", "message_text": None, "receipt_ref": None,
            "csv_file_path": None, "feedback": None}
    return asyncio.run(run_agent({**base, **state}))


def fake_llm(monkeypatch, extraction=None, category=CAT, fail_extraction=False):
    async def fake_extract(image_bytes):
        if fail_extraction:
            raise LLMUnavailableError("all models down")
        return extraction, LLMCallInfo(provider="test", model="fake")

    async def fake_categorize(data):
        return (CategorizationResult(category=category, item="Team Dinner", goods_or_service="Service"),
                LLMCallInfo(provider="test", model="fake"))

    monkeypatch.setattr(graph_module.extractor, "extract", fake_extract)
    monkeypatch.setattr(graph_module.categorizer, "categorize", fake_categorize)


def send_receipt(monkeypatch, **extraction_kw):
    data = dict(vendor="ABC Restaurant", expense_date=date.today(), expense_amount=2450, currency="INR",
                description="Team dinner")
    data.update(extraction_kw)
    fake_llm(monkeypatch, ReceiptExtraction(**data))
    ref = asyncio.run(storage_service.save_file(png(), ".png"))
    return agent(input_type="receipt", receipt_ref=ref)


def press(action, expense_id, value=None):
    return agent(input_type="feedback", feedback={"action": action, "expense_id": expense_id, "value": value})


# ---------------------------------------------------------------- human-in-the-loop
def test_receipt_creates_pending_card_with_confirm_and_correct_buttons(monkeypatch):
    out = send_receipt(monkeypatch)
    assert "Vendor: ABC Restaurant" in out["response_text"] and "Expense: ₹2,450" in out["response_text"]
    labels = [b[0] for row in out["buttons"] for b in row]
    assert labels == ["✅ Confirm", "✏️ Correct"]
    assert expense_service.get(out["expense_id"], USER).status == "pending"


def test_pending_expenses_are_not_counted_in_reports(monkeypatch):
    send_receipt(monkeypatch)
    assert tools.get_expense_summary(None)["empty"] is True


def test_confirm_with_corrections_stores_ai_and_final_values(monkeypatch):
    out = send_receipt(monkeypatch, expense_amount=12500)
    eid = out["expense_id"]
    press("set_category", eid, str(CAT_IDX_FIXED))                  # human picks a different category
    agent(input_type="feedback", feedback={"action": "field", "expense_id": eid, "value": "expense_amount"})
    agent(input_type="text", message_text="₹12,050")                # typed correction
    done = press("confirm", eid)
    assert "Expense saved successfully" in done["response_text"]

    e = expense_service.get(eid, USER)
    assert (e.status, e.category_ai, e.category, e.corrected_by_user) == ("confirmed", CAT, CAT_FIXED, True)
    assert float(e.expense_amount_ai) == 12500 and float(e.expense_amount) == 12050
    assert sorted(e.corrected_fields) == ["category", "expense_amount"]
    assert e.final_result["category"] == CAT_FIXED and e.ai_extraction["extraction"]["expense_amount"] == 12500
    with SessionLocal() as s:
        rows = {c.field: (c.ai_value, c.final_value) for c in s.query(ExpenseCorrection).all()}
    assert rows["category"] == (CAT, CAT_FIXED)
    assert rows["expense_amount"] == ("12500.00", "12050.00")


def test_confirm_without_changes_marks_not_corrected_and_double_tap_is_safe(monkeypatch):
    eid = send_receipt(monkeypatch)["expense_id"]
    press("confirm", eid)
    assert expense_service.get(eid, USER).corrected_by_user is False
    again = press("confirm", eid)
    assert again["toast"] == "That expense was already handled."
    with SessionLocal() as s:
        assert s.query(ExpenseCorrection).count() == 0


def test_invalid_typed_amount_is_rejected_and_state_kept(monkeypatch):
    eid = send_receipt(monkeypatch)["expense_id"]
    press("field", eid, "expense_amount")
    bad = agent(input_type="text", message_text="lots")
    assert "positive number" in bad["response_text"]
    ok = agent(input_type="text", message_text="3000")
    assert "Expense: ₹3,000" in ok["response_text"]


def test_cancel_leaves_edit_mode(monkeypatch):
    eid = send_receipt(monkeypatch)["expense_id"]
    press("field", eid, "vendor")
    agent(input_type="text", message_text="cancel")
    parsed(monkeypatch, intent="unknown")
    out = agent(input_type="text", message_text="hello")           # now treated as a normal message
    assert out["response_text"].startswith("I can help with")


def test_user_cannot_touch_someone_elses_expense(monkeypatch):
    eid = send_receipt(monkeypatch)["expense_id"]
    out = agent(user_id=222, input_type="feedback", feedback={"action": "confirm", "expense_id": eid})
    assert out["toast"] == "I can't find that expense."
    assert expense_service.get(eid, USER).status == "pending"


def test_extraction_failure_still_allows_manual_entry(monkeypatch):
    fake_llm(monkeypatch, fail_extraction=True)
    ref = asyncio.run(storage_service.save_file(png(), ".png"))
    out = agent(input_type="receipt", receipt_ref=ref)
    assert "couldn't read this receipt" in out["response_text"] and out["buttons"]
    eid = out["expense_id"]
    blocked = press("confirm", eid)
    assert "Missing" in blocked["toast"]
    press("set_category", eid, "0")
    for field, value in (("expense_amount", "500"), ("expense_date", "08/10/2026")):
        press("field", eid, field)
        agent(input_type="text", message_text=value)
    assert "saved successfully" in press("confirm", eid)["response_text"]


def test_discard(monkeypatch):
    eid = send_receipt(monkeypatch)["expense_id"]
    press("discard", eid)
    assert expense_service.get(eid, USER).status == "discarded"


def test_unreadable_image_gives_friendly_error(monkeypatch):
    fake_llm(monkeypatch, ReceiptExtraction())

    async def boom(image_bytes):
        raise ValueError("Not a readable image")
    monkeypatch.setattr(graph_module.extractor, "extract", boom)
    ref = asyncio.run(storage_service.save_file(b"junk", ".jpg"))
    assert "readable image" in agent(input_type="receipt", receipt_ref=ref)["error"]


# ---------------------------------------------------------------- reports / charts / excel / receipts
def confirmed(monkeypatch, **kw):
    out = send_receipt(monkeypatch, **kw)
    press("confirm", out["expense_id"])
    return out["expense_id"]


def parsed(monkeypatch, **request):
    async def fake_parse(text, today):
        return ReportRequest(**request), "llm"
    monkeypatch.setattr(graph_module.parser, "parse", fake_parse)


def test_category_report_returns_text_and_chart_from_the_database(monkeypatch):
    confirmed(monkeypatch)
    confirmed(monkeypatch, vendor="Rail Co", expense_amount=1000)
    parsed(monkeypatch, intent="category_breakdown", period=Period(kind="this_month"))
    out = agent(message_text="Show me expenses by category")
    assert f"{CAT} — ₹3,450 (100%)" in out["response_text"]
    assert len(out["chart_paths"]) == 1 and os.path.getsize(out["chart_paths"][0]) > 1000


def test_monthly_trend_and_last_six_months(monkeypatch):
    confirmed(monkeypatch)
    parsed(monkeypatch, intent="monthly_trend", period=Period(kind="last_n_months", months=6))
    out = agent(message_text="Show me the last 6 months of expenses")
    assert out["response_text"].count("—") >= 6 and out["chart_paths"]


def test_excel_report_contains_the_filtered_rows(monkeypatch):
    confirmed(monkeypatch)
    parsed(monkeypatch, intent="excel_report", period=Period(kind="this_month"))
    out = agent(message_text="Download this month's Excel report")
    wb = load_workbook(out["excel_file_path"])
    assert wb.sheetnames == ["Summary", "Expenses", "Category Summary", "Monthly Summary"]
    rows = [r for r in wb["Expenses"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert len(rows) == 1 and rows[0][2] == "ABC Restaurant" and rows[0][7] == 2450


def test_excel_category_filter(monkeypatch):
    confirmed(monkeypatch)
    parsed(monkeypatch, intent="excel_report", period=Period(kind="this_month"), category="Travel")
    out = agent(message_text="Excel for Travel")
    assert out["excel_file_path"] is None and "No confirmed expenses" in out["response_text"]


def test_receipt_retrieval_by_amount_and_keyword_returns_the_stored_image(monkeypatch):
    confirmed(monkeypatch)
    parsed(monkeypatch, intent="receipt_lookup", amount=2450, search_text="restaurant", period=Period(kind="all"))
    out = agent(message_text="Show me the receipt for the ₹2,450 restaurant expense")
    assert len(out["receipt_images"]) == 1 and os.path.isfile(out["receipt_images"][0]["path"])
    assert "ABC Restaurant" in out["receipt_images"][0]["caption"]


def test_receipt_retrieval_by_category(monkeypatch):
    confirmed(monkeypatch)
    parsed(monkeypatch, intent="receipt_lookup", category="Travel", period=Period(kind="all"))
    assert "couldn't find" in agent(message_text="Show me Travel receipts")["response_text"]


def test_mixed_currencies_are_never_added_together(monkeypatch):
    confirmed(monkeypatch)
    confirmed(monkeypatch, currency="USD", expense_amount=50)
    result = tools.get_expense_summary(None)
    assert result["currency"] == "INR" and result["total"] == 2450 and result["other_currency_count"] == 1


def test_unknown_request_shows_help(monkeypatch):
    parsed(monkeypatch, intent="unknown")
    assert agent(message_text="what's the weather")["response_text"].startswith("I can help with")


# ---------------------------------------------------------------- telegram layer
def test_unauthorized_user_is_rejected(monkeypatch):
    from unittest.mock import AsyncMock
    from app.telegram import bot as bot_module
    send = AsyncMock()
    monkeypatch.setattr(bot_module.bot_instance, "send_message", send)
    update = {"update_id": 9001, "message": {"chat": {"id": 999}, "from": {"id": 999}, "text": "hi"}}
    asyncio.run(bot_module.handle_telegram_update(update))
    (chat_id, text), _ = send.call_args
    assert chat_id == 999 and "not authorized" in text and "999" in text     # the reply tells them their id


def test_duplicate_updates_are_processed_once(monkeypatch):
    from unittest.mock import AsyncMock
    from app.telegram import bot as bot_module
    send = AsyncMock()
    monkeypatch.setattr(bot_module.bot_instance, "send_message", send)
    update = {"update_id": 9002, "message": {"chat": {"id": 1}, "from": {"id": 1}, "text": "hi"}}
    asyncio.run(bot_module.handle_telegram_update(update))
    asyncio.run(bot_module.handle_telegram_update(update))
    assert send.call_count == 1


def test_webhook_acknowledges_immediately(monkeypatch):
    from fastapi.testclient import TestClient
    import app.main as main
    import app.telegram.bot as bot_module
    seen = []
    monkeypatch.setattr(bot_module, "schedule_update", lambda data: seen.append(data))
    resp = TestClient(main.app).post("/telegram/webhook", json={"update_id": 1})
    assert resp.status_code == 200 and seen == [{"update_id": 1}]


# ---------------------------------------------------------------- feature: category bar chart & timeline
def test_category_bar_chart_on_request(monkeypatch):
    confirmed(monkeypatch)
    parsed(monkeypatch, intent="chart", chart_type="category_bar", period=Period(kind="this_month"))
    out = agent(message_text="Show me a bar graph of this month's expenses by category")
    assert len(out["chart_paths"]) == 1 and os.path.getsize(out["chart_paths"][0]) > 1000
    assert f"{CAT} — ₹2,450" in out["response_text"]


def test_timeline_graph_uses_real_dates_from_the_database(monkeypatch):
    first = date.today().replace(day=1)
    confirmed(monkeypatch, expense_date=first, expense_amount=1000)
    confirmed(monkeypatch, expense_date=date.today(), expense_amount=500, vendor="Rail Co")
    parsed(monkeypatch, intent="chart", chart_type="timeline", period=Period(kind="this_month"),
           granularity="day", cumulative=True)
    out = agent(message_text="Show me expenses over time")
    assert len(out["chart_paths"]) == 1 and os.path.getsize(out["chart_paths"][0]) > 1000
    assert "Cumulative expenses over time" in out["response_text"] and "₹1,500" in out["response_text"]


def test_timeline_direct_tool_returns_continuous_cumulative_series(monkeypatch):
    first = date.today().replace(day=1)
    confirmed(monkeypatch, expense_date=first, expense_amount=1000)
    res = tools.generate_timeline_chart(None, first.isoformat(), date.today().isoformat(), None, "day", True)
    assert res["granularity"] == "day" and res["rows"][0]["cumulative"] == 1000
    assert res["rows"][-1]["cumulative"] == 1000 and len(res["rows"]) == date.today().day


def test_timeline_with_no_data_says_so(monkeypatch):
    parsed(monkeypatch, intent="chart", chart_type="timeline", period=Period(kind="this_month"))
    out = agent(message_text="expenses over time")
    assert out["chart_paths"] == [] and "No confirmed expenses" in out["response_text"]


# ---------------------------------------------------------------- restored CSV import, /add wizard, duplicates
def test_csv_import_saves_confirmed_rows_and_reports_skips(tmp_path):
    path = tmp_path / "expenses.csv"
    path.write_text("date,vendor,amount,currency,category,description\n"
                    "04/10/2026,Lavish Restaurant,454,INR,zakat,lunch\n"
                    "2026-10-05,Rail Co,\"1,200.50\",INR,not-a-category,ticket\n"
                    "bad-date,Nobody,10,INR,,\n", encoding="utf-8")
    out = agent(input_type="csv", csv_file_path=str(path))
    assert "Imported 2 expense(s)" in out["response_text"] and "Skipped 1" in out["response_text"]
    assert not path.exists()                                              # temp file removed
    with SessionLocal() as s:
        rows = {e.vendor: e for e in s.query(Expense).all()}
    assert rows["Lavish Restaurant"].category == "zakath" and rows["Lavish Restaurant"].source == "csv"
    assert rows["Rail Co"].category == "Miscellaneous" and float(rows["Rail Co"].expense_amount) == 1200.5
    assert tools.get_expense_summary(None)["count"] == 2


def add_event(**event):
    return agent(input_type="manual_add", manual_event=event)


def test_add_wizard_end_to_end_through_the_graph_and_database(monkeypatch):
    add_event(kind="start")
    # while the wizard is open, plain text is an answer - NOT a report request (the parser must not run)
    async def must_not_run(text, today):
        raise AssertionError("text was routed to the report parser during /add")
    monkeypatch.setattr(graph_module.parser, "parse", must_not_run)

    asked = agent(message_text="750")
    assert "choose a category" in asked["response_text"] and asked["buttons"][0][0][1] == "ma:0"
    picked = add_event(kind="category", value="3")
    assert ALL_CATEGORIES[3].value in picked["response_text"] and picked["edit_message"] is True
    done = agent(message_text="Recharge for the office phone")
    assert "added successfully" in done["response_text"] and "₹750" in done["response_text"]

    with SessionLocal() as s:
        e = s.query(Expense).one()
    assert (e.status, e.source, e.category, float(e.expense_amount)) == ("confirmed", "manual", ALL_CATEGORIES[3].value, 750.0)
    assert e.expense_date == date.today() and e.description == "Recharge for the office phone"
    assert tools.get_expense_summary(None)["total"] == 750.0            # counted in reports immediately
    # the wizard is closed: the next text goes back to normal report parsing
    parsed(monkeypatch, intent="unknown")
    assert agent(message_text="hello")["response_text"].startswith("I can help with")


def test_cancel_closes_the_add_wizard(monkeypatch):
    add_event(kind="start")
    out = agent(input_type="feedback", feedback={"action": "cancel_edit"})
    assert out["response_text"] == "Okay, cancelled." and expense_service.get_manual(USER) is None


def test_starting_a_field_edit_closes_the_add_wizard(monkeypatch):
    eid = send_receipt(monkeypatch)["expense_id"]
    add_event(kind="start")
    press("field", eid, "vendor")
    assert expense_service.get_manual(USER) is None and expense_service.get_edit(USER) == (eid, "vendor")


def test_same_receipt_uploaded_twice_is_flagged_as_duplicate(monkeypatch):
    fake_llm(monkeypatch, ReceiptExtraction(vendor="ABC Restaurant", expense_date=date.today(), expense_amount=2450,
                                            currency="INR"))
    first = agent(input_type="receipt", receipt_ref=asyncio.run(storage_service.save_file(png(), ".png")))
    second = agent(input_type="receipt", receipt_ref=asyncio.run(storage_service.save_file(png(), ".png")))
    assert "already uploaded" not in first["response_text"]
    assert "already uploaded" in second["response_text"] and f"#{first['expense_id']}" in second["response_text"]
    assert first["expense_id"] != second["expense_id"]          # warn, never block: the user decides


def test_category_prompt_lists_every_category_and_description():
    from app.ai.receipt_extractor import _categorizer_system
    prompt = _categorizer_system()
    for c in ALL_CATEGORIES:
        assert c.value in prompt


def test_report_parser_prompt_uses_a_real_category_in_its_example():
    from app.ai.receipt_extractor import _parser_system
    prompt = _parser_system(date.today())
    assert f'"category": "{ALL_CATEGORIES[2].value}"' in prompt and '"category": "Travel"' not in prompt


def test_legacy_category_remap_plan():
    from scripts.remap_categories import LEGACY_MAP, plan
    assert all(v in ALL_CATEGORIES for v in LEGACY_MAP.values())
    changes, unmapped = plan({"Travel": 2, "Food & Meals": 3, ALL_CATEGORIES[9].value: 4, "Mystery": 1, None: 5})
    assert changes == [("Travel", ALL_CATEGORIES[2].value, 2), ("Food & Meals", ALL_CATEGORIES[6].value, 3)]
    assert unmapped == [("Mystery", 1)]                      # current categories and NULLs are left alone


def test_telegram_add_command_and_category_button(monkeypatch):
    """The real bot path: /add -> amount -> tap a category button (callback 'ma:<index>')."""
    from unittest.mock import AsyncMock
    from app.telegram import bot as bot_module
    send, edit, answer = AsyncMock(), AsyncMock(), AsyncMock()
    monkeypatch.setattr(bot_module.bot_instance, "send_message", send)
    monkeypatch.setattr(bot_module.bot_instance, "edit_message", edit)
    monkeypatch.setattr(bot_module.bot_instance, "answer_callback", answer)

    def say(update_id, text):
        asyncio.run(bot_module.handle_telegram_update(
            {"update_id": update_id, "message": {"chat": {"id": 111}, "from": {"id": 111}, "text": text}}))

    say(7001, "/add")
    assert "Send the amount" in send.call_args[0][1]
    say(7002, "300")
    assert "choose a category" in send.call_args[0][1] and len(send.call_args[0][2][0]) == 2   # buttons attached

    asyncio.run(bot_module.handle_telegram_update({"update_id": 7003, "callback_query": {
        "id": "cb1", "from": {"id": 111}, "data": "ma:2",
        "message": {"chat": {"id": 111}, "message_id": 55}}}))
    edit.assert_called_once()
    assert edit.call_args[0][:2] == (111, 55) and ALL_CATEGORIES[2].value in edit.call_args[0][2]
    answer.assert_called_once_with("cb1", ALL_CATEGORIES[2].value)

    say(7004, "bus ticket")
    assert "added successfully" in send.call_args[0][1]
    assert tools.get_expense_summary(None)["total"] == 300.0


def test_stray_add_button_without_a_draft_is_harmless(monkeypatch):
    from unittest.mock import AsyncMock
    from app.telegram import bot as bot_module
    answer = AsyncMock()
    monkeypatch.setattr(bot_module.bot_instance, "send_message", AsyncMock())
    monkeypatch.setattr(bot_module.bot_instance, "edit_message", AsyncMock())
    monkeypatch.setattr(bot_module.bot_instance, "answer_callback", answer)
    asyncio.run(bot_module.handle_telegram_update({"update_id": 7010, "callback_query": {
        "id": "cb2", "from": {"id": 111}, "data": "ma:1", "message": {"chat": {"id": 111}, "message_id": 9}}}))
    answer.assert_called_once_with("cb2", "That entry expired. Send /add to start again.")


def test_database_url_is_masked_and_connections_time_out():
    from app.config.settings import settings
    from app.database.models import connect_args_for, describe_db_url
    shown = describe_db_url("postgresql://postgres:s3cret@db.example.com:5433/expense_db")
    assert "s3cret" not in shown and "db.example.com:5433" in shown
    assert connect_args_for("sqlite:///x.db") == {}
    assert connect_args_for("postgresql://u:p@h/d") == {"connect_timeout": settings.DB_CONNECT_TIMEOUT_SECONDS}
    assert settings.DB_CONNECT_TIMEOUT_SECONDS <= 10          # an unreachable host must fail in seconds, not minutes
